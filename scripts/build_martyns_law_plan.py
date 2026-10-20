#!/usr/bin/env python3
"""Build Martyn's Law portal planning brief (DOCX)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

OUT = Path(__file__).resolve().parents[1] / "docs" / "Martyns-Law-Portal-Planning.docx"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
NAVY = RGBColor(0x0B, 0x25, 0x45)
GRAY = RGBColor(0x55, 0x55, 0x55)
BLACK = RGBColor(0x00, 0x00, 0x00)
HEADER_FILL = "F2F4F7"
CALLOUT_FILL = "EEF3F9"


def set_run(run, *, size=11, bold=False, color=BLACK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_cell_shading(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_table_geometry(table, widths_in: list[float], indent_dxa: int = 120):
    total = int(sum(w * 1440 for w in widths_in))
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    def upsert(tag, attrs):
        el = tblPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tblPr.append(el)
        for k, v in attrs.items():
            el.set(qn(k), str(v))

    upsert("w:tblW", {"w:w": total, "w:type": "dxa"})
    upsert("w:tblInd", {"w:w": indent_dxa, "w:type": "dxa"})

    # layout fixed
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    # grid
    existing = tbl.find(qn("w:tblGrid"))
    if existing is not None:
        tbl.remove(existing)
    grid = OxmlElement("w:tblGrid")
    for w in widths_in:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(w * 1440)))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(widths_in[i] * 1440)))
            tcW.set(qn("w:type"), "dxa")
            for m_tag, val in (
                ("w:top", 80),
                ("w:bottom", 80),
                ("w:start", 120),
                ("w:end", 120),
            ):
                # ensure mar
                pass
            mar = tcPr.find(qn("w:tcMar"))
            if mar is None:
                mar = OxmlElement("w:tcMar")
                tcPr.append(mar)
            for edge, val in (
                ("top", 80),
                ("left", 120),
                ("bottom", 80),
                ("right", 120),
            ):
                node = mar.find(qn(f"w:{edge}"))
                if node is None:
                    node = OxmlElement(f"w:{edge}")
                    mar.append(node)
                node.set(qn("w:w"), str(val))
                node.set(qn("w:type"), "dxa")


def style_table(table, widths_in: list[float], header=True):
    set_table_geometry(table, widths_in)
    table.style = "Table Grid"
    if header:
        for cell in table.rows[0].cells:
            set_cell_shading(cell, HEADER_FILL)
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run(r, size=10, bold=True, color=NAVY)
    for row in table.rows[1 if header else 0 :]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    set_run(r, size=10, color=BLACK)


def add_title_block(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("PROJECT PLANNING BRIEF")
    set_run(r, size=10, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Martyn’s Law Portal")
    set_run(r, size=22, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(
        "Planning document for a mosque-facing portal that helps venues determine "
        "Martyn’s Law (Terrorism (Protection of Premises) Act) duties, build public "
        "protection procedures/plans, train people, and evidence readiness."
    )
    set_run(r, size=11, color=GRAY)


def add_h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_run(r, size=16, bold=True, color=BLUE)
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_h2(doc, text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        set_run(r, size=13, bold=True, color=BLUE)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run(r, size=11, color=BLACK)
    return p


def add_bullets(doc, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        r = p.add_run(item)
        set_run(r, size=11, color=BLACK)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15


def add_kv_table(doc, rows: list[tuple[str, str]], widths=(1.875, 4.625)):
    t = doc.add_table(rows=len(rows), cols=2)
    for i, (a, b) in enumerate(rows):
        t.rows[i].cells[0].text = a
        t.rows[i].cells[1].text = b
        set_cell_shading(t.rows[i].cells[0], HEADER_FILL)
        for cell in t.rows[i].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run(r, size=10, bold=(cell == t.rows[i].cells[0]), color=BLACK)
    style_table(t, list(widths), header=False)
    doc.add_paragraph()


def add_matrix(doc, headers: list[str], rows: list[list[str]], widths: list[float]):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri + 1].cells[ci].text = val
    style_table(t, widths, header=True)
    doc.add_paragraph()


def add_callout(doc, label: str, text: str):
    t = doc.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    set_cell_shading(cell, CALLOUT_FILL)
    cell.text = ""
    p = cell.paragraphs[0]
    r1 = p.add_run(f"{label}  ")
    set_run(r1, size=10, bold=True, color=NAVY)
    r2 = p.add_run(text)
    set_run(r2, size=10, color=BLACK)
    style_table(t, [6.5], header=False)
    doc.add_paragraph()


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for level, size, before, after in (
        (1, 16, 16, 8),
        (2, 13, 12, 6),
        (3, 12, 8, 4),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE if level < 3 else RGBColor(0x1F, 0x4D, 0x78)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_footer(doc: Document):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Martyn’s Law Portal · Planning brief · Confidential working draft")
    set_run(r, size=8, color=GRAY)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top", "bottom", "left", "right"):
        setattr(section, f"{side}_margin", Inches(1))
    configure_styles(doc)
    add_footer(doc)

    add_title_block(doc)

    add_kv_table(
        doc,
        [
            ("Document type", "Planning brief / product & architecture scope"),
            ("Audience", "Programme leads, trustees, product, content & engineering"),
            ("Date", "31 July 2026"),
            ("Status", "Working draft for review"),
        ],
    )

    add_callout(
        doc,
        "Purpose in one line",
        "Help mosques and other faith venues understand whether Martyn’s Law applies, "
        "which tier they fall into, what procedures or plans they must have, train the "
        "right people, and keep an evidence trail for trustees and regulators.",
    )

    add_h1(doc, "1. Objectives")
    add_body(
        doc,
        "The portal should make Martyn’s Law practical for mosque trustees, premises/"
        "security leads, event organisers and front-of-house volunteers. It is not a "
        "generic compliance checklist; it links capacity/tier determination to public "
        "protection procedures, training and improvement actions.",
    )
    add_bullets(
        doc,
        [
            "Determine whether a premises or qualifying event is in scope, and whether Standard or Enhanced duties apply.",
            "Guide the Responsible Person through public protection procedures (and, where required, a public protection plan).",
            "Convert gaps into owned actions with due dates and evidence notes.",
            "Assign role-based awareness and procedure drills for staff and volunteers.",
            "Create a durable evidence trail for trustees, insurers and regulatory readiness.",
            "Reuse the Beacon Mosque pack/token assessment pattern already used for safeguarding.",
        ],
    )

    add_h1(doc, "2. Problem statement")
    add_body(
        doc,
        "Martyn’s Law (Terrorism (Protection of Premises) Act) places clear duties on "
        "those responsible for certain premises and events. Many mosques already think "
        "about protective security, but practice is uneven: capacity thresholds are not "
        "documented, procedures for evacuate / invacuate / lockdown / communicate are "
        "oral, and large gatherings (Jummah, Eid, funerals, conferences) are treated as "
        "routine rather than qualifying events.",
    )
    add_body(
        doc,
        "Without a shared digital workspace, readiness stays informal, evidence is "
        "scattered across WhatsApp and filing cabinets, and trustees cannot see whether "
        "the mosque is Standard-tier ready, Enhanced-tier obligated, or still unset.",
    )

    add_h1(doc, "3. Legislative framing (portal wording)")
    add_body(
        doc,
        "Portal content must stay aligned with official Home Office / regulator guidance "
        "as it is published. The product should encode the duty model, not invent "
        "legal advice.",
    )
    add_matrix(
        doc,
        ["Concept", "Portal treatment"],
        [
            [
                "Qualifying premises",
                "Guided questionnaire on use, public access and capacity thresholds",
            ],
            [
                "Qualifying events",
                "Event packs for large gatherings that may trigger duties even if the site itself does not",
            ],
            [
                "Standard duty",
                "Public protection procedures + people aware of their roles",
            ],
            [
                "Enhanced duty",
                "Procedures plus documented public protection plan and appropriate security measures",
            ],
            [
                "Responsible Person",
                "Named in pack; owns procedures/plan and review cycle",
            ],
            [
                "Regulatory readiness",
                "Evidence export for trustee files; not a substitute for legal advice",
            ],
        ],
        [2.0, 4.5],
    )

    add_h1(doc, "4. Users and roles")
    add_matrix(
        doc,
        ["Role", "Primary job in the portal"],
        [
            [
                "Responsible Person",
                "Confirm premises/event details, complete procedures/plan, own actions",
            ],
            [
                "Trustees / chair",
                "Approve tier determination, review residual risk, resource actions",
            ],
            [
                "Premises & security lead",
                "Complete site, access, CCTV, entry control and protective measures",
            ],
            [
                "Operations / events lead",
                "Register qualifying events; brief stewards; run drills",
            ],
            [
                "Front-of-house / stewards",
                "Complete awareness modules; acknowledge procedure roles",
            ],
            [
                "Beacon Mosque programme team",
                "Issue packs, monitor progress, support mosques, report aggregate readiness",
            ],
        ],
        [2.25, 4.25],
    )

    add_h1(doc, "5. Scope")
    add_h2(doc, "5.1 In scope")
    add_bullets(
        doc,
        [
            "Mosque / venue packs with secure access links (token model).",
            "Capacity and use questionnaire with Standard / Enhanced / Out-of-scope outcome.",
            "Public protection procedures builder (evacuate, invacuate, lockdown, communicate with emergency services).",
            "Enhanced-tier public protection plan sections and security-measure checklist.",
            "Qualifying event packs linked to a parent premises pack.",
            "RAG-rated readiness domains, evidence notes and open actions.",
            "Role-based awareness training catalogue with completion status.",
            "Admin view for Beacon Mosque to create packs and monitor progress.",
            "Trustee summary export (printable / PDF in a later phase).",
        ],
    )
    add_h2(doc, "5.2 Out of scope (initial release)")
    add_bullets(
        doc,
        [
            "Live liaison with police CTSA systems or SIA casework.",
            "Purchase or configuration of physical security equipment.",
            "Detailed threat intelligence feeds.",
            "Full multi-site enterprise portfolio management beyond mosque packs.",
            "Legal determination letters (the portal supports readiness; it does not certify compliance).",
            "Multi-language content beyond English in phase 1.",
        ],
    )

    add_h1(doc, "6. Assessment framework")
    add_body(
        doc,
        "Assessment is organised by duty layers and operational domains. Each domain has "
        "a short checklist, RAG rating, evidence notes and improvement actions. Domain "
        "wording should be reviewed against final statutory guidance before pilot launch.",
    )
    add_matrix(
        doc,
        ["Layer", "Domains", "Focus"],
        [
            [
                "Duty determination",
                "Capacity & use; Qualifying events; Responsible Person",
                "Is the Act in play, and who owns it?",
            ],
            [
                "Standard procedures",
                "Evacuate; Invacuate; Lockdown; Communicate",
                "Public protection procedures everyone can follow",
            ],
            [
                "Enhanced plan",
                "Risk assessment; Security measures; Plan document; Review cycle",
                "Documented plan and proportionate measures",
            ],
            [
                "People readiness",
                "Role briefings; Awareness training; Drills & rehearsals",
                "Right people know what to do",
            ],
            [
                "Assurance",
                "Evidence pack; Trustee oversight; Annual refresh",
                "Board can see residual risk",
            ],
        ],
        [1.7, 2.6, 2.2],
    )

    add_h2(doc, "6.1 Core checklist (phase 1)")
    add_matrix(
        doc,
        ["Domain", "Checklist items"],
        [
            [
                "Capacity & use",
                "Peak public capacity recorded; use categories confirmed; threshold rationale saved",
            ],
            [
                "Responsible Person",
                "Named individual; contact details; deputy identified; trustee endorsement",
            ],
            [
                "Evacuate",
                "Routes; assembly points; roles; accessibility considerations; practice date",
            ],
            [
                "Invacuate / lockdown",
                "Safe areas; how to secure; communication cascade; visitor handling",
            ],
            [
                "Communicate",
                "Who dials 999; what information is given; internal alert method; media holding line",
            ],
            [
                "Enhanced plan",
                "Written plan; security measures proportionate to risk; review date; owner",
            ],
            [
                "Training & drills",
                "Induction for stewards; annual refresh; drill log; lessons learned",
            ],
        ],
        [1.9, 4.6],
    )

    add_h2(doc, "6.2 RAG meaning")
    add_matrix(
        doc,
        ["RAG", "Meaning"],
        [
            ["Green", "Procedure/plan is in place, evidenced and reviewed."],
            ["Amber", "Partial practice — gaps known and actions planned."],
            ["Red", "Missing or unsafe — urgent action required."],
            ["Unset", "Not yet reviewed."],
        ],
        [1.5, 5.0],
    )

    add_h1(doc, "7. Architecture")
    add_body(
        doc,
        "The Martyn’s Law portal should reuse the Beacon Mosque self-assessment stack "
        "(Next.js App Router, Cloudflare-friendly APIs, D1/SQLite pack storage, secure "
        "access tokens) and extend it with a tier-determination engine and procedure/"
        "plan builders.",
    )

    add_h2(doc, "7.1 Logical architecture")
    add_matrix(
        doc,
        ["Layer", "Components", "Responsibility"],
        [
            [
                "Experience",
                "Pack home, tier wizard, procedure builder, event packs, training view, trustee summary, admin console",
                "Mosque-facing UX and programme operations",
            ],
            [
                "Application APIs",
                "/api/martyns/packs, /domains, /procedures, /events, /actions, /training",
                "CRUD, validation, recommendations",
            ],
            [
                "Domain services",
                "Tier engine; procedure templates; action recommender; training mapper; export builder",
                "Encode Martyn’s Law product rules",
            ],
            [
                "Persistence",
                "D1/SQLite tables for packs, domains, procedures, events, actions, completions",
                "Durable evidence and progress",
            ],
            [
                "Shared platform",
                "Auth/token access, CMS/admin identity, file/evidence store (phase 2), observability",
                "Security and operations",
            ],
        ],
        [1.5, 2.7, 2.3],
    )

    add_h2(doc, "7.2 System context")
    add_body(
        doc,
        "External actors: Mosque users (Responsible Person, trustees, stewards), Beacon "
        "Mosque programme admins, and later optional content/guidance sources. The portal "
        "does not directly integrate with police systems in phase 1; it prepares evidence "
        "and procedure packs that humans can share with CTSA / local partners when needed.",
    )
    add_bullets(
        doc,
        [
            "Browser clients → Next.js routes/components.",
            "API handlers validate pack tokens and write to D1.",
            "Catalog modules (procedure templates, training map) live in versioned code/content.",
            "Admin console lists packs, RAG summaries and open actions for programme support.",
            "Optional future: PDF export worker; evidence object storage; SSO for multi-user packs.",
        ],
    )

    add_h2(doc, "7.3 Suggested data model (phase 1)")
    add_matrix(
        doc,
        ["Entity", "Key fields", "Notes"],
        [
            [
                "martyns_packs",
                "id, access_token, venue_name, responsible_person, contact, status, tier_outcome, created_at",
                "Mirrors assessment_packs pattern",
            ],
            [
                "martyns_domains",
                "pack_id, domain_id, rag, checklist_json, evidence_notes, updated_at",
                "RAG + checklist completion",
            ],
            [
                "martyns_procedures",
                "pack_id, procedure_type, steps_json, owners_json, last_drill_at, status",
                "evacuate / invacuate / lockdown / communicate",
            ],
            [
                "martyns_events",
                "id, pack_id, name, expected_capacity, date_label, tier_flags, status",
                "Qualifying event overlays",
            ],
            [
                "martyns_actions",
                "id, pack_id, domain_id, title, owner_role, due_label, status",
                "Improvement backlog",
            ],
            [
                "martyns_training",
                "id, pack_id, person_label, role, module_id, completed_at",
                "Awareness completions",
            ],
        ],
        [1.7, 2.7, 2.1],
    )

    add_h2(doc, "7.4 Duty & assurance flow")
    add_bullets(
        doc,
        [
            "1. Create pack → capture venue/event profile.",
            "2. Tier engine → Out of scope / Standard / Enhanced (with rationale stored).",
            "3. Procedure builder → four public protection procedures completed and owned.",
            "4. If Enhanced → public protection plan + security measures checklist.",
            "5. People readiness → assign modules and record drills.",
            "6. RAG summary → actions → trustee review → monitoring / annual refresh.",
        ],
    )

    add_h2(doc, "7.5 Relationship to safeguarding portal")
    add_body(
        doc,
        "Keep Martyn’s Law as a sibling product domain to the safeguarding self-assessment, "
        "not a buried subsection. Share platform primitives (pack tokens, RAG, actions, "
        "admin list) but keep catalogues, copy and legal framing separate so neither "
        "duty is diluted. Cross-links from Premises & Security / Hate Crime domains in "
        "the safeguarding pack can deep-link into the Martyn’s Law pack for the same mosque.",
    )

    add_h1(doc, "8. Training pathway design")
    add_body(
        doc,
        "Training follows the assessment. Amber/red procedure domains recommend the "
        "modules that close the gap. Completions feed the People readiness layer.",
    )
    add_matrix(
        doc,
        ["Module", "Audience", "Outcome"],
        [
            [
                "Martyn’s Law awareness",
                "All staff & stewards",
                "Understand why procedures exist and personal role",
            ],
            [
                "Responsible Person briefing",
                "Named RP / deputies",
                "Own tier outcome, procedures/plan and review cycle",
            ],
            [
                "Trustee oversight",
                "Trustees",
                "Ask the right readiness questions and resource gaps",
            ],
            [
                "Procedure drill facilitation",
                "Operations / security leads",
                "Run short rehearsals and capture lessons learned",
            ],
            [
                "Large-event stewardship",
                "Event stewards",
                "Entry flow, observation, escalation and calm communication",
            ],
            [
                "Annual refresher",
                "All relevant roles",
                "Keep knowledge current on a scheduled cycle",
            ],
        ],
        [2.1, 1.8, 2.6],
    )

    add_h1(doc, "9. End-to-end user journey")
    add_bullets(
        doc,
        [
            "Beacon Mosque creates a Martyn’s Law pack and shares a secure link.",
            "Mosque confirms venue details and names the Responsible Person.",
            "Tier wizard records capacity/use and returns Standard, Enhanced or Out of scope.",
            "Leads complete procedure domains (and plan domains if Enhanced).",
            "Portal recommends actions and training from amber/red domains.",
            "Qualifying events are registered against the pack before large gatherings.",
            "Trustees review summary; open actions are owned with due dates.",
            "Stewards complete awareness modules; drills update readiness.",
            "Pack moves draft → in review → monitoring; annual refresh begins.",
        ],
    )

    add_h1(doc, "10. Product surfaces")
    add_matrix(
        doc,
        ["Surface", "Purpose"],
        [
            ["Pack home", "Tier outcome, RAG overview, open actions, training progress"],
            ["Tier wizard", "Capacity/use questionnaire with stored rationale"],
            ["Procedure builder", "Step-by-step evacuate / invacuate / lockdown / communicate"],
            ["Enhanced plan workspace", "Plan sections + security measures checklist"],
            ["Event pack pages", "Large gathering readiness linked to parent venue"],
            ["Training / learner view", "Assigned modules, status, refreshers"],
            ["Trustee summary", "High-level readiness and priority risks"],
            ["Admin / programme console", "Create packs, support mosques, export metrics"],
        ],
        [2.2, 4.3],
    )

    add_h1(doc, "11. Delivery phases")
    add_matrix(
        doc,
        ["Phase", "When", "Outcome"],
        [
            [
                "0 — Align",
                "Week 1–2",
                "Confirm legal content owners, pilot mosques, success metrics",
            ],
            [
                "1 — Tier + procedures MVP",
                "Week 3–9",
                "Packs, tier wizard, four procedures, RAG, actions, admin list",
            ],
            [
                "2 — Enhanced + events",
                "Week 10–15",
                "Public protection plan workspace; qualifying event packs",
            ],
            [
                "3 — Training & report",
                "Week 16–20",
                "Awareness modules, recommendations, trustee summary export",
            ],
            [
                "4 — Scale",
                "Week 21+",
                "Evidence uploads, refreshers, multi-language, broader rollout",
            ],
        ],
        [2.0, 1.4, 3.1],
    )

    add_h2(doc, "11.1 Pilot success criteria")
    add_bullets(
        doc,
        [
            "At least 5 pilot mosques complete tier determination with stored rationale.",
            "Each Standard-tier pilot has all four procedures rated (not unset) after walkthrough.",
            "Enhanced-tier pilots (if any) have a draft public protection plan section set.",
            "80% of named Responsible Persons finish the RP briefing within 4 weeks.",
            "Trustees can state their tier outcome and top 3 residual risks from the summary.",
        ],
    )

    add_h1(doc, "12. Data, privacy and duty boundaries")
    add_bullets(
        doc,
        [
            "Store venue contacts, assessment answers, procedure text, action logs and training completions only.",
            "Do not store sensitive personal intelligence, watchlists or individual threat reports.",
            "Use secure pack tokens / authenticated access; do not publish mosque scores publicly.",
            "Define retention periods for packs, procedure versions and drill logs.",
            "Provide a clear privacy notice for mosque users and programme administrators.",
            "Label content as readiness support, not legal advice or regulatory certification.",
        ],
    )

    add_h1(doc, "13. Risks and mitigations")
    add_matrix(
        doc,
        ["Risk", "Mitigation"],
        [
            [
                "Portal mistaken for legal certification",
                "Onboarding disclaimer; trustee summary language reviewed by content owner",
            ],
            [
                "Guidance changes after build",
                "Versioned catalogue; named legal/content owner; change log",
            ],
            [
                "Low completion by busy volunteers",
                "Short modules; mobile-friendly builders; facilitator support",
            ],
            [
                "Inconsistent tier self-assessment",
                "Clear threshold prompts; rationale required; admin spot-checks",
            ],
            [
                "Overlap confusion with safeguarding pack",
                "Sibling navigation; shared mosque identity; distinct catalogues",
            ],
            [
                "Access link sharing too widely",
                "Token rotation; contact verification; update audit trail",
            ],
        ],
        [2.4, 4.1],
    )

    add_h1(doc, "14. Open decisions")
    add_matrix(
        doc,
        ["Decision", "Options / notes"],
        [
            [
                "Product packaging",
                "Standalone /martyns routes vs tab inside a unified mosque readiness hub",
            ],
            [
                "Identity model",
                "Magic-link packs vs accounts for Responsible Person + stewards",
            ],
            [
                "Tier engine strictness",
                "Advisory questionnaire vs hard-gated workflow before procedures unlock",
            ],
            [
                "Evidence uploads",
                "Notes only in MVP vs file uploads (plans, site maps, drill logs)",
            ],
            [
                "Event model",
                "Child packs vs lightweight event records on the parent pack",
            ],
            [
                "Content authority",
                "Beacon Mosque editorial vs external CT/security partner review",
            ],
        ],
        [2.1, 4.4],
    )

    add_h1(doc, "15. Workstreams and owners")
    add_matrix(
        doc,
        ["Workstream", "Owner", "Key deliverables"],
        [
            ["Product / UX", "TBC", "Tier wizard, procedure builder, event + trustee views"],
            ["Duty content", "TBC", "Checklist wording, procedure templates, disclaimers"],
            ["Training content", "TBC", "Awareness + RP + trustee modules"],
            ["Engineering", "TBC", "Packs API, tier fields, procedures schema, admin"],
            ["Pilot support", "TBC", "Mosque onboarding, walkthroughs, feedback loop"],
            ["Governance", "TBC", "Privacy, retention, legal review cadence"],
        ],
        [1.7, 1.2, 3.6],
    )

    add_h1(doc, "16. Immediate next steps")
    add_bullets(
        doc,
        [
            "Confirm pilot mosques and named Responsible Persons.",
            "Approve phase-1 checklist and procedure template wording against current guidance.",
            "Decide packaging: sibling of safeguarding portal vs unified readiness hub.",
            "Extend data model (martyns_* tables) beside existing assessment_* schema.",
            "Prototype tier wizard + four procedure domains on the current pack UX.",
            "Schedule fortnightly steering check-ins through phase 2.",
        ],
    )

    add_h1(doc, "17. Document control")
    add_matrix(
        doc,
        ["Version", "Date", "Notes"],
        [["0.1", "31 July 2026", "Initial planning brief with architecture for review"]],
        [1.2, 1.5, 3.8],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
